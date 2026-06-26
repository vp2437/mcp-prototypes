from openpyxl import load_workbook, Workbook
from mcp.server.fastmcp import FastMCP
import os

mcp = FastMCP("Excel Server")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))


@mcp.tool()
def server_directory():
    """
    Return the directory containing server.py.
    """
    return BASE_DIR

@mcp.tool()
def list_excel_files():
    """
    List Excel files located beside server.py.
    """
    return [f for f in os.listdir(BASE_DIR) if f.endswith(".xlsx")]

@mcp.tool()
def create_excel(file_path: str, headers: list) -> str:
    """
    Create a new Excel workbook with the specified column headers.

    Args:
        file_path: Name or path of the Excel file to create.
        headers: List of column names for the first row.

    Returns:
        Confirmation message.
    """
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    full_path = os.path.join(BASE_DIR, file_path)
    wb.save(full_path)
    return f"Created {file_path} with headers: {headers}"

@mcp.tool()
def read_excel(file_path: str) -> str:
    """
    Read all rows from an Excel workbook.

    Use this tool when a user asks to:
    - view a spreadsheet
    - inspect spreadsheet contents
    - summarize spreadsheet data
    - read an Excel file

    Args:
        file_path: Excel filename such as 'sample.xlsx'.

    Returns:
        Worksheet contents as tab-separated text.
    """
    full_path = os.path.join(BASE_DIR, file_path)
    wb = load_workbook(full_path)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append("\t".join(str(cell) if cell is not None else "" for cell in row))
    return "\n".join(rows)

@mcp.tool()
def read_range(file_path: str, range_str: str) -> str:
    """
    Read a specific range of cells from the active worksheet.

    Example:
        range_str = "A1:C5"

    Args:
        file_path: Path to the Excel file.
        range_str: Excel range.

    Returns:
        Range contents as tab-separated text.
    """
    full_path = os.path.join(BASE_DIR, file_path)
    wb = load_workbook(full_path)
    ws = wb.active
    rows = []
    for row in ws[range_str]:
        rows.append("\t".join(str(cell.value) if cell.value is not None else "" for cell in row))
    return "\n".join(rows)

@mcp.tool()
def read_sheet(file_path: str, sheet_name: str) -> str:
    """
    Read all data from a specified worksheet.

    Args:
        file_path: Path to the Excel file.
        sheet_name: Name of the worksheet.

    Returns:
        Worksheet contents as tab-separated text.
    """
    full_path = os.path.join(BASE_DIR, file_path)
    wb = load_workbook(full_path)
    ws = wb[sheet_name]
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append("\t".join(str(cell) if cell is not None else "" for cell in row))
    return "\n".join(rows)

@mcp.tool()
def write_row(file_path: str, row_data: list) -> str:
    """
    Append a new row to the active worksheet.

    Args:
        file_path: Path to the Excel file.
        row_data: List of values to append.

    Returns:
        Confirmation message.
    """
    full_path = os.path.join(BASE_DIR, file_path)
    wb = load_workbook(full_path)
    ws = wb.active
    ws.append(row_data)
    wb.save(full_path)
    return f"Row added to {file_path}"

@mcp.tool()
def update_cell(file_path: str, cell: str, value: str) -> str:
    """
    Update a single cell in the active worksheet.

    Example:
        cell = "B2"

    Args:
        file_path: Path to the Excel file.
        cell: Cell reference.
        value: New value.

    Returns:
        Confirmation message.
    """
    full_path = os.path.join(BASE_DIR, file_path)
    wb = load_workbook(full_path)
    ws = wb.active
    ws[cell] = value
    wb.save(full_path)
    return f"Wrote '{value}' to {cell} in {file_path}"

if __name__ == "__main__":
    mcp.run(transport="stdio")